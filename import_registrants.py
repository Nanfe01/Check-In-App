"""Import registrants and generate pass links.

    python import_registrants.py registrants.xlsx --base-url https://checkin.example.org

Reads a .xlsx or .csv with a "Full Name" (or "Name") column and an "email" column.
Writes <input>_with_links.<ext> containing your original columns plus `pass_link`.

Uses the database in DATABASE_URL if set (e.g. your Neon database), otherwise a local checkin.db.
Safe to re-run with an updated sheet: people already imported (matched by email)
keep their existing token, so links you already sent keep working.
"""
import argparse
import csv
import sys
from pathlib import Path

import store

NAME_HEADERS = ("full name", "fullname", "name")
EMAIL_HEADERS = ("email", "e-mail", "email address")


def read_rows(path: Path):
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = [["" if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, newline="", encoding="utf-8-sig") as f:
            rows = [[c.strip() for c in r] for r in csv.reader(f)]
    rows = [r for r in rows if any(r)]
    if not rows:
        sys.exit("The file is empty.")
    return rows[0], rows[1:]


def find_col(headers, options):
    lowered = [h.lower().strip() for h in headers]
    for opt in options:
        if opt in lowered:
            return lowered.index(opt)
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--base-url", required=True, help="Public URL of the server, e.g. https://checkin.example.org")
    args = ap.parse_args()

    path = Path(args.file)
    headers, data = read_rows(path)
    name_i, email_i = find_col(headers, NAME_HEADERS), find_col(headers, EMAIL_HEADERS)
    if name_i is None or email_i is None:
        sys.exit(f"Need a 'Full Name' or 'Name' column and an 'email' column. Found: {headers}")

    base = args.base_url.rstrip("/")
    store.init_db()
    print(f"Writing to: {store.backend_name()}")
    out_rows, added, kept, skipped = [], 0, 0, 0
    with store.connect() as conn:
        for r in data:
            r = r + [""] * (len(headers) - len(r))
            name, email = r[name_i], r[email_i].lower()
            if not name:
                skipped += 1
                out_rows.append(r + ["(skipped: no name)"])
                continue
            row = conn.execute("SELECT * FROM registrants WHERE lower(email)=?", (email,)).fetchone() if email else None
            if row is None:
                row = store.add_registrant(conn, name, email)
                added += 1
            else:
                kept += 1
            out_rows.append(r + [f"{base}/pass/{row['token']}"])

    out = path.with_name(f"{path.stem}_with_links{path.suffix}")
    out_headers = headers + ["pass_link"]
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(out_headers)
        for r in out_rows:
            ws.append(r)
        ws.column_dimensions[ws.cell(1, len(out_headers)).column_letter].width = 60
        out = out.with_suffix(".xlsx")
        wb.save(out)
    else:
        with open(out, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f); w.writerow(out_headers); w.writerows(out_rows)

    print(f"New: {added}  Already imported (token kept): {kept}  Skipped: {skipped}")
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
